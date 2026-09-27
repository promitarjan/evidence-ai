import io
from pathlib import Path
from docx import Document
from models.schemas import ExportDocxRequest, AcceptedCitation, DocumentParagraph
from services.docx_exporter import generate_citation_docx

def test_docx_export_and_formatting():
    test_dir = Path(__file__).parent / "test_data"
    original_docx = test_dir / "sample_paper.docx"

    with open(original_docx, "rb") as f:
        orig_bytes = f.read()

    request = ExportDocxRequest(
        originalFileName="sample_paper.docx",
        documentTitle="Climate Impacts on Rice Agriculture in South Asia",
        paragraphs=[
            DocumentParagraph(index=0, text="Rice is an essential staple crop that sustains over two billion people across Asia. Bangladesh is one of the world's most climate-vulnerable countries due to its low elevation and river delta topography."),
            DocumentParagraph(index=1, text="Over the last three decades, extreme temperature variability has significantly disrupted national crop yields. Recent field trials indicate that high temperature stress during the flowering stage reduces rice grain fertility by as much as 25%."),
            DocumentParagraph(index=2, text="Rice accounts for approximately 70% of agricultural land use in Bangladesh. Despite rapid technological modernization in irrigation systems, rural farmer adaptation remains constrained by seasonal flood risks.")
        ],
        acceptedCitations=[
            AcceptedCitation(
                claimId="c-1",
                paragraphIdx=1,
                sentenceText="Recent field trials indicate that high temperature stress during the flowering stage reduces rice grain fertility by as much as 25%.",
                formattedCitation="(Rahman, 2025)",
                formattedBibEntry="Rahman, M. et al. (2025). Thermal stress and yield variability in South Asian river deltas. Journal of Agricultural Climate Science, 18(2), 104-118."
            ),
            AcceptedCitation(
                claimId="c-2",
                paragraphIdx=0,
                sentenceText="Bangladesh is one of the world's most climate-vulnerable countries due to its low elevation and river delta topography.",
                formattedCitation="(IPCC, 2023)",
                formattedBibEntry="IPCC. (2023). Climate Change 2023: Synthesis Report. Contribution of Working Groups I, II and III."
            )
        ]
    )

    docx_stream = generate_citation_docx(request, orig_bytes)
    assert docx_stream is not None
    assert docx_stream.getbuffer().nbytes > 1000

    # Reload the generated document to verify structure
    exported_doc = Document(docx_stream)
    
    full_text = "\n".join(p.text for p in exported_doc.paragraphs)
    print("--- Exported Document Paragraphs Preview ---")
    for i, p in enumerate(exported_doc.paragraphs):
        print(f"  P[{i}]: {p.text[:75]}...")

    # Verification assertions:
    assert "(Rahman, 2025)" in full_text, "In-text citation (Rahman, 2025) missing from exported document"
    assert "(IPCC, 2023)" in full_text, "In-text citation (IPCC, 2023) missing from exported document"
    assert "References" in full_text, "References section heading missing"
    assert "Rahman, M. et al. (2025)" in full_text, "Bibliography entry for Rahman missing"
    assert "IPCC. (2023)" in full_text, "Bibliography entry for IPCC missing"

    # Save to test_data for inspection
    output_path = test_dir / "exported_verified_test.docx"
    docx_stream.seek(0)
    with open(output_path, "wb") as f_out:
        f_out.write(docx_stream.read())
    print(f"\nSaved test DOCX to: {output_path}")
    print("\n✓ ROUND-TRIP DOCX EXPORT AND FORMATTING MECHANICAL TEST PASSED!")

if __name__ == "__main__":
    test_docx_export_and_formatting()
